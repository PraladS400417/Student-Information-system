from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

OUT='Pralad_Raya_S400417_Individual_Test_Report.docx'
d=Document(); sec=d.sections[0]; sec.top_margin=Inches(.75); sec.bottom_margin=Inches(.7); sec.left_margin=Inches(.85); sec.right_margin=Inches(.85)
styles=d.styles; styles['Normal'].font.name='Aptos'; styles['Normal'].font.size=Pt(10.5); styles['Normal'].paragraph_format.space_after=Pt(7)
for name,size in [('Title',22),('Heading 1',15),('Heading 2',11.5)]:
 s=styles[name]; s.font.name='Aptos Display' if name=='Title' else 'Aptos'; s.font.size=Pt(size); s.font.color.rgb=RGBColor(0,0,0); s.font.bold=True
 s.paragraph_format.space_before=Pt(14 if name!='Title' else 0); s.paragraph_format.space_after=Pt(6)

def p(x='',style=None): return d.add_paragraph(x,style)
def h(x,level=1): d.add_heading(x,level)
def tbl(headers,rows,widths=None):
 t=d.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.style='Table Grid'
 for i,x in enumerate(headers): t.rows[0].cells[i].text=x
 for row in rows:
  cells=t.add_row().cells
  for i,x in enumerate(row): cells[i].text=str(x)
 for row in t.rows:
  for c in row.cells:
   tcPr=c._tc.get_or_add_tcPr(); mar=OxmlElement('w:tcMar')
   for side in ('top','left','bottom','right'):
    el=OxmlElement('w:'+side); el.set(qn('w:w'),'90'); el.set(qn('w:type'),'dxa'); mar.append(el)
   tcPr.append(mar)
   for pp in c.paragraphs:
    pp.paragraph_format.space_after=Pt(0)
    for r in pp.runs: r.font.size=Pt(9)
 for c in t.rows[0].cells:
  shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),'D9E6F2'); c._tc.get_or_add_tcPr().append(shd)
  for r in c.paragraphs[0].runs: r.bold=True
 d.add_paragraph('')
 return t

d.add_paragraph('Individual Software Testing Report','Title')
p('Student Record System  |  PRT582 Software Group Project')
tbl(['Student','Student ID','Group name and number'],[['Pralad Raya','S400417','To be supplied by the group']])
p('Report date: 4 October 2026. Evidence basis: project source code and SQL dump, assignment brief, and a prepared Robot Framework suite. Runtime testing is pending because PHP and MySQL were unavailable in the inspected environment.')
h('Contents')
for x in ['1 Individual Contribution Statement','2 Individual Non Functional Testing','3 Individual Functional Testing','4 Ethical Consideration','5 Conclusion and Reflection','Appendix A Evidence and Reproduction','Appendix B Robot Framework Execution']: p(x)

h('1 Individual Contribution Statement')
p('My assigned area is the teacher workflow: authentication, PIN verification, profile, password change, language course enrollment, student or course entry, and enrollment history. I inspected the corresponding PHP pages and database schema, defined acceptance scenarios, and prepared seven Robot Framework browser tests in pralad_teacher_tests.robot. I did not claim browser execution or performance measurements because this environment did not contain PHP, MySQL, or XAMPP.')
h('Project background and scope',2)
p('The application is a PHP and MySQL student information system with teacher and admin portals. The teacher side uses PHP sessions and mysqli queries. The supplied SQL dump is sqlfile/acourse.sql and the database configured by both portals is Acourse. The README names a different database and a teacher login absent from the dump, so setup must follow the actual configuration and seed data. My scope is the teacher interface; other group members cover admin configuration, teacher management, and grades.')
h('Requirements and user stories',2)
tbl(['ID','User story','Acceptance criterion'],[
('F1','As a teacher, I can sign in and out.','Valid credentials open the teacher area; invalid credentials stay on login; logout removes access.'),
('F2','As a teacher, I can verify my PIN.','Wrong PIN shows an error; correct PIN opens language course enrollment.'),
('F3','As a teacher, I can view and update my profile.','My registration number is displayed; permitted changes persist after reload.'),
('F4','As a teacher, I can manage enrollment.','Selected course and academic options save once and appear in the appropriate history.'),
('F5','As a teacher, I can change my password.','Wrong current password is rejected; a successful change requires the new password on next login.'),
('NF1','As a teacher, I can use the forms reliably.','Labels, errors, keyboard order, and responses are understandable.'),
('NF2','As an operator, I can trust security and response time.','Unauthorised requests are rejected and routine pages meet an agreed timing target.')])
h('Test plan and automation strategy',2)
p('Use a disposable local Acourse database imported from sqlfile/acourse.sql. Start PHP and MySQL, then the PHP server at localhost:8000. Use the seeded teacher 070101 / 1234 and PIN 185748 only in an isolated local test copy. Run read-only browser tests first, then manual create, edit, and password tests with a disposable account. Reset the database between destructive runs. Capture Robot output.xml, log.html, report.html, screenshots, timestamps, PHP version, and database version. A test is complete only when its expected and actual outcomes are recorded.')
h('Automation tool comparison',2)
tbl(['Tool','Strength','Limit','Decision'],[
('Robot Framework with SeleniumLibrary','Readable cases and built-in HTML reports; required by brief','Browser and driver setup required','Selected for teacher acceptance flows'),
('Playwright','Reliable modern browser automation','Different reporting and syntax; not the named assignment tool','Useful comparison only'),
('PHPUnit','Good for isolated PHP units','Application has no existing unit-test seams or Composer setup','Not primary for UI acceptance')])

h('2 Individual Non Functional Testing')
h('Usability Testing',2)
p('Method: ask 3–5 representative users to complete teacher login, find PIN verification, locate enrollment history, and recover from a wrong PIN. Record task completion, time, wrong turns, and a 1–5 clarity rating. The current evidence is a source inspection, not a user study. The menu labels mix Teacher and Student terminology (for example, “Enroll Student” and “student Marks” in includes/menubar.php), and my-profile.php displays “Teacher Registration” for the profile screen. These are credible usability concerns to validate with users. Improvement: make role and action labels consistent, and show field-specific errors beside the relevant input.')
h('Performance Testing',2)
p('Method: after the local server works, use Apache JMeter to test GET /index.php, authenticated profile, and enrollment history with 1, 5, then 20 concurrent users for 2 minutes each. Record sample count, median, 95th percentile, error rate, CPU, and database load. Proposed acceptance targets for local baseline are 95th percentile below 2 seconds and error rate below 1%; these are project targets, not measured results. No timings are reported because the application could not be launched here. Recheck any slow database query with EXPLAIN and repeat under the same environment before drawing conclusions.')
h('Security Testing',2)
p('Distinct security tool: OWASP ZAP passive scan on localhost after login, followed by manual checks of unauthorised page access, wrong PIN, tampered read-only form values, and a safe SQL injection probe in a disposable database. Do not scan public hosts or real user data. Static inspection found SQL queries that concatenate request values in index.php, pincode-verification.php, my-profile.php, and enroll.php, creating injection risk. Passwords are compared as MD5 hashes, which is unsuitable for current password storage. The profile upload moves the supplied filename into staffphoto without an explicit type or filename validation in my-profile.php. These are code findings; exploitability and ZAP alerts are not claimed as observed test results. Recommended fixes: prepared statements, password_hash/password_verify, strict upload validation, CSRF protection, and server-side authorization checks.')

h('3 Individual Functional Testing')
h('Manual acceptance scenarios',2)
tbl(['ID','Steps and expected result','Status'],[
('M01','Sign in with 070101 / 1234; teacher page opens. Retry with wrong password; error appears.','Pending runtime'),
('M02','Open enroll.php without a session; redirect to login.','Pending runtime'),
('M03','Submit wrong PIN then 185748; error then enrollment form.','Pending runtime'),
('M04','Edit teacher name, reload profile; new name persists without debug output.','Pending runtime'),
('M05','Change password with wrong then correct current password; verify next login.','Pending runtime'),
('M06','Submit course enrollment, refresh, check exactly one saved record and history display.','Pending runtime'),
('M07','Use keyboard only through login, PIN, enrollment, and profile forms; note focus and errors.','Pending runtime')])
h('Robot Framework automated testing execution and results',2)
p('The accompanying pralad_teacher_tests.robot implements seven browser checks: valid login, invalid login, protected-page redirect, wrong PIN, correct PIN, profile registration number, and enrollment history load. Execution result: not run. No pass or fail counts, screenshots, or Robot output files exist yet. This is a prepared suite, not evidence of successful execution. Installation and command are in Appendix B. The team should run it after local setup and replace this paragraph with the dated Robot summary and attached log files.')
h('Source findings relevant to acceptance',2)
tbl(['Finding','Evidence','Likely effect and follow-up'],[
('F01 History mismatch','enroll.php writes courseenrolls; enroll-history.php queries student.','Language course enrollment may not appear in this history. Confirm UI intent and test both histories.'),
('F02 Debug output in profile','my-profile.php calls var_dump($_POST) on submit.','Form data may be printed. Remove debug output and retest.'),
('F03 Teacher field trust','enroll.php accepts posted staffregno and PIN into insert query.','Tampering may misattribute enrollment. Derive identity from session and validate.'),
('F04 Weak query handling','Several pages concatenate raw form values into SQL.','Use prepared statements, then repeat negative and security tests.')])

h('4 Ethical Consideration')
p('I would run tests only on an authorised local copy with synthetic or supplied sample records. I would avoid posting screenshots that expose passwords, PINs, session cookies, or personal information. Destructive tests need a disposable account and a database backup. Automated security checks should be limited to localhost and recorded accurately. Reporting a code risk is different from claiming a successful exploit or completed test, so every result in this report is labelled by its evidence type.')

h('5 Conclusion and Reflection')
p('I produced a teacher workflow test plan and a Robot Framework suite grounded in the actual PHP pages. The source review identified a likely enrollment history inconsistency, debug output, and input-handling risks. The work is not a completed runtime test campaign: the browser, usability, load, and security tool runs remain to be performed after PHP and MySQL are available. The most useful next step is to start the local site, execute the suite, collect its generated reports, and update this document with observed outcomes and defect reproduction evidence.')
p('My technical learning from this contribution is the need to trace each UI action to its database table before writing acceptance checks. I also learned to separate requirements, expected outcomes, static findings, and observed results. The main challenge was the mismatch between README instructions and the actual configuration and SQL dump; resolving it required checking both config files and the seed records.')
h('Further work',2)
p('Complete the manual cases, user study, JMeter baseline, and ZAP scan; fix confirmed defects; rerun regression cases; then integrate verified results into the group report. Group name, group number, and any video link must be supplied by the group before final submission.')

h('Appendix A Evidence and Reproduction')
p('Repository files inspected: README.md; sqlfile/acourse.sql; includes/config.php; admin/includes/config.php; index.php; includes/menubar.php; pincode-verification.php; enroll.php; enroll-history.php; my-profile.php; change-password.php; course.php. Assignment source: PRT582 Software Group Project.docx, Individual Submission Format and Robot Framework requirement. Setup guide: RUN_AND_TEST_GUIDE.txt in the project folder. No live run, screenshots, or logs were available at the time of this report.')
h('Appendix B Robot Framework Execution')
p('Install Python 3, Google Chrome, Robot Framework, and SeleniumLibrary on the test machine. Start MySQL and the PHP server, confirm http://localhost:8000/, then run:')
p('python -m pip install robotframework robotframework-seleniumlibrary')
p('python -m robot --outputdir pralad_robot_results pralad_teacher_tests.robot')
p('The suite defaults to localhost:8000, Chrome, and the seeded teacher credentials. Override variables with --variable BASE_URL:<url> and other --variable options as needed. Save output.xml, log.html, and report.html with the final submission. Use a local disposable database and change seeded passwords after testing.')
p('Video presentation link: to be added by the group. Detailed non-functional data and logs: pending execution.')

foot=sec.footer.paragraphs[0]; foot.alignment=WD_ALIGN_PARAGRAPH.CENTER; foot.text='Pralad Raya  |  S400417  |  PRT582'
d.save(OUT)
print(OUT)
